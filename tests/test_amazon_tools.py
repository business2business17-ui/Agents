"""Regression tests for the Amazon agents (Product Intelligence, Feed Compiler, Feed Error Agent): shared workbook tools and scripts.

Run:  python3 -m unittest discover -s tests -v      (needs openpyxl, lxml)
Builds a synthetic Amazon-like workbook (Template with example row 6, list validation via a defined name, hidden
sheet, merged cells, fake macro part) and checks the safety guarantees of patch / guard / dry-run / report parsing.
"""
import csv
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SH = os.path.join(ROOT, "shared", "amazon")
PI = os.path.join(ROOT, "plugins", "amazon-product-intelligence", "skills", "amazon-product-intelligence", "scripts")
ER = os.path.join(ROOT, "plugins", "amazon-feed-error-agent", "skills", "amazon-feed-error-agent", "scripts")


def run(script, *args, ok=(0,)):
    p = subprocess.run([sys.executable, script, *map(str, args)], capture_output=True, text=True)
    assert p.returncode in ok, f"{os.path.basename(script)} {args} -> {p.returncode}\n{p.stdout}\n{p.stderr}"
    return p


def make_feed(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template"
    ws["A1"] = "TemplateType=x"
    ws.merge_cells("A1:C1")
    keys = ["feed_product_type", "item_sku", "external_product_id", "item_name", "color_name", "generic_keywords"]
    for i, k in enumerate(keys, 1):
        ws.cell(5, i, k)
        ws.cell(6, i, "example")
    dv = DataValidation(type="list", formula1="=ValidColors", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("E7:E200")
    dd = wb.create_sheet("Data Definitions")
    dd["A1"] = "Field"
    vv = wb.create_sheet("Valid Values")
    vv["A1"], vv["A2"] = "Blue", "Red"
    vv.sheet_state = "hidden"
    wb.defined_names["ValidColors"] = DefinedName("ValidColors", attr_text="'Valid Values'!$A$1:$A$2")
    tmp = path + ".tmp"
    wb.save(tmp)
    with zipfile.ZipFile(tmp) as zi, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zo:
        for n in zi.namelist():
            zo.writestr(n, zi.read(n))
        zo.writestr("xl/vbaProject.bin", b"\x00FAKEVBA")
    os.remove(tmp)


class AmazonTools(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = tempfile.mkdtemp()
        cls.feed = os.path.join(cls.d, "clean.xlsm")
        make_feed(cls.feed)

    def p(self, name):
        return os.path.join(self.d, name)

    def write_json(self, name, data):
        with open(self.p(name), "w", encoding="utf-8") as f:
            json.dump(data, f)
        return self.p(name)

    def test_pricing_policy_v3(self):
        run(os.path.join(SH, "pricing_engine.py"), "--self-test")
        out = json.loads(run(os.path.join(SH, "pricing_engine.py"), "--sale", "24.99", "--marketplace", "DE").stdout)
        self.assertEqual((out["standard_price"], out["business_price"]), (27.77, 24.99))

    def test_user_decided_tiers_and_b2b_min(self):
        out = json.loads(run(os.path.join(SH, "pricing_engine.py"), "--sale", "24.99", "--marketplace", "DE", "--tiers",
                             "2:5,4:10", "--tier-basis", "business", "--b2b-min", "deepest-tier").stdout)
        self.assertEqual(out["business_min_price"], out["quantity_tiers"][-1]["price"])
        self.assertEqual(out["guardrails"]["policy_status"], "USER_DECISION")
        nobasis = run(os.path.join(SH, "pricing_engine.py"), "--sale", "24.99", "--marketplace", "DE", "--tiers", "2:5", ok=(1,))
        self.assertIn("PRICE_DATA_REQUIRED", nobasis.stdout)

    def test_margin_calc_proposes_ladder_within_margin(self):
        args = ["--cost", "9", "--referral-pct", "15", "--fba-fee", "3.2", "--channel", "fba", "--vat-pct", "19",
                "--target-margin", "15", "--basis-price", "24.99", "--quantities", "2,4", "--json"]
        out = json.loads(run(os.path.join(SH, "margin_calc.py"), *args).stdout)["fba"]
        ladder = out["proposed_ladder"]["tiers"]
        self.assertEqual([t["quantity"] for t in ladder], [2, 4])
        for t in ladder:  # every proposed tier keeps the target margin
            self.assertGreaterEqual(float(t["margin"].rstrip("%")), 15.0)
        p = run(os.path.join(SH, "margin_calc.py"), "--cost", "9", "--referral-pct", "15", "--mfn-fee", "4.5", "--channel", "mfn",
                "--vat-pct", "19", "--target-margin", "25", "--price", "24.99", ok=(1,))
        self.assertIn("BELOW TARGET", p.stdout)

    def test_performance_metrics(self):
        base = ["--cost", "9", "--referral-pct", "15", "--fba-fee", "3.2", "--channel", "fba", "--vat-pct", "19",
                "--lines", "24.99:300,22.49:120", "--total-sales", "10195.80", "--json"]
        out = json.loads(run(os.path.join(SH, "performance_calc.py"), *base, "--ad-spend", "650", "--ad-sales", "2100",
                             "--target-margin", "10", ok=(0, 1)).stdout)
        self.assertEqual(out["gmv_gross"], "10195.80")
        self.assertEqual(out["acos"], "31.0%")
        self.assertEqual(out["tacos"], "6.4%")
        self.assertEqual(out["roas"], "3.23")
        self.assertEqual(out["net_profit"], "1264.53")
        p = run(os.path.join(SH, "performance_calc.py"), *base, "--ad-spend", "650", "--ad-sales", "2100",
                "--target-margin", "15", ok=(1,))
        self.assertIn("below target", p.stdout)

    def test_sales_basis_has_no_silent_default(self):
        args = ["--cost", "9", "--referral-pct", "15", "--fba-fee", "3.2", "--channel", "fba", "--vat-pct", "19",
                "--lines", "24.99:300", "--ad-spend", "100", "--ad-sales", "500"]
        p = run(os.path.join(SH, "performance_calc.py"), *args, ok=(1,))
        self.assertEqual(p.returncode, 1)
        self.assertIn("sales basis unknown", p.stderr)
        net = json.loads(run(os.path.join(SH, "performance_calc.py"), *args, "--total-sales", "6300", "--json", ok=(0, 1)).stdout)
        self.assertTrue(net["assumptions"]["sales_basis"].startswith("net (DETECTED"))

    def test_b2b_max_is_tied_to_sale_price(self):
        out = json.loads(run(os.path.join(SH, "pricing_engine.py"), "--sale", "24.99", "--marketplace", "DE", "--b2b-max-pct", "20").stdout)
        self.assertEqual(out["business_max_price"], 29.99)
        self.assertGreaterEqual(out["business_max_price"], out["sale_price"])
        self.assertIn("max(Business Price, Sale Price)", out["guardrails"]["business_max_rule"])

    def test_patch_guard_roundtrip(self):
        cells = self.write_json("c.json", [
            {"cell": "B7", "value": "SKU-1"}, {"cell": "C7", "value": "0012345678905", "type": "text"},
            {"cell": "D7", "value": "=evil()"}, {"cell": "E7", "value": "Blue"}])
        out = self.p("out.xlsm")
        run(os.path.join(SH, "xlsm_patch.py"), "--source", self.feed, "--cells", cells, "--out", out, "--sheet", "Template")
        run(os.path.join(SH, "workbook_guard.py"), "--source", self.feed, "--output", out, "--sheet", "Template", "--approved", cells)
        wb = openpyxl.load_workbook(out)
        ws = wb["Template"]
        self.assertEqual(ws["C7"].value, "0012345678905")
        self.assertEqual(ws["C7"].data_type, "s")
        self.assertEqual(ws["D7"].data_type, "s")  # text, not a formula
        self.assertEqual(ws["B6"].value, "example")
        with zipfile.ZipFile(out) as z:
            self.assertIn("xl/vbaProject.bin", z.namelist())

    def test_patch_refuses_protected_rows_and_overwrites(self):
        bad = self.write_json("bad.json", [{"cell": "B6", "value": "x"}])
        p = run(os.path.join(SH, "xlsm_patch.py"), "--source", self.feed, "--cells", bad, "--out", self.p("o1.xlsm"),
                "--sheet", "Template", ok=(1,))
        self.assertIn("READ-ONLY", p.stderr)
        self.assertFalse(os.path.exists(self.p("o1.xlsm")))
        same = run(os.path.join(SH, "xlsm_patch.py"), "--source", self.feed, "--cells", bad, "--out", self.feed,
                   "--sheet", "Template", ok=(1,))
        self.assertIn("differ", same.stderr)

    def test_guard_detects_tampering(self):
        cells = self.write_json("c2.json", [{"cell": "B7", "value": "SKU-2"}])
        out = self.p("out2.xlsm")
        run(os.path.join(SH, "xlsm_patch.py"), "--source", self.feed, "--cells", cells, "--out", out, "--sheet", "Template")
        wb = openpyxl.load_workbook(out)
        wb["Template"]["B6"] = "HACK"
        wb["Data Definitions"]["A1"] = "Changed"
        bad = self.p("tampered.xlsx")
        wb.save(bad)
        p = run(os.path.join(SH, "workbook_guard.py"), "--source", self.feed, "--output", bad, "--sheet", "Template",
                "--approved", cells, ok=(1,))
        self.assertIn("NOT READY", p.stdout)

    def test_dry_run_catches_enum_and_identifier_problems(self):
        cells = self.write_json("c3.json", [
            {"cell": "E7", "value": "blue"}, {"cell": "E8", "value": "Green"},
            {"cell": "C7", "value": 4006381333931, "type": "number"}, {"cell": "B7", "value": "S1"}, {"cell": "B8", "value": "S1"}])
        p = run(os.path.join(SH, "validate_cells.py"), "--source", self.feed, "--cells", cells, "--sheet", "Template",
                "--key-column", "B", ok=(1,))
        for code in ("ENUM_AMBIGUOUS", "ENUM_INVALID", "IDENTIFIER_NOT_TEXT", "DUPLICATE_FEED_ROW"):
            self.assertIn(code, p.stdout)

    def test_inspect_confirms_row7_boundary(self):
        run(os.path.join(SH, "xlsm_inspect.py"), self.feed, "--json", self.p("insp.json"))
        d = json.load(open(self.p("insp.json"), encoding="utf-8"))
        self.assertTrue(d["macros_present"])
        self.assertEqual(d["data_start_check"]["status"], "DATA_START_CONFIRMED")

    def test_processing_report_and_comparison(self):
        wb = openpyxl.load_workbook(self.feed)
        ws = wb["Template"]
        ws["B7"], ws["E7"] = "SKU-1", "Bluee"
        ws["E7"].fill = PatternFill("solid", fgColor="FFFFA500")
        s = wb.create_sheet("Feed Processing Summary", 0)
        s.append(["Error code", "Category of error", "Store", "Error message", "Affected field", "Impacted column", "Number of errors"])
        s.append([8058, "Error", "DE", "Invalid value", "color_name", "E", 1])
        rep = self.p("rep.xlsx")
        wb.save(rep)
        run(os.path.join(ER, "parse_processing_report.py"), rep, "--marketplace", "DE", "--json", self.p("p1.json"))
        d = json.load(open(self.p("p1.json"), encoding="utf-8"))
        self.assertEqual(d["totals"]["errors"], 1)
        self.assertEqual(d["findings"][0]["error_code"], "8058")
        d2 = dict(d, findings=[])
        self.write_json("p2.json", d2)
        p = run(os.path.join(ER, "compare_reports.py"), self.p("p1.json"), self.p("p2.json"))
        self.assertIn("Resolved: 1", p.stdout)

    def test_gtin_and_handoff(self):
        p = run(os.path.join(PI, "gtin_check.py"), "4006381333931", "4006381333932", ok=(1,))
        self.assertIn("GTIN_VALID", p.stdout)
        self.assertIn("check digit should be 1", p.stdout)
        ex = run(os.path.join(PI, "handoff_tool.py"), "--example").stdout
        path = self.p("h.json")
        open(path, "w").write(ex)
        run(os.path.join(PI, "handoff_tool.py"), "validate", path)
        rec = json.loads(ex)
        rec["pricing"]["standard_price"] = 1.0
        self.write_json("h_bad.json", rec)
        run(os.path.join(PI, "handoff_tool.py"), "validate", self.p("h_bad.json"), ok=(1,))

    def test_seo_import_tiers_and_marketplace_isolation(self):
        import csv as _csv
        de = self.p("cerebro_de.csv")
        with open(de, "w", newline="", encoding="utf-8-sig") as f:
            w = _csv.writer(f, delimiter=";")
            w.writerow(["Keyword Phrase", "Search Volume", "Cerebro IQ Score"])
            for r in [("handyhülle pixel 8", "12.400", 1), ("pixel 8 hülle", "9.800", 1), ("hülle pixel 8", "9.800", 1),
                      ("samsung galaxy hülle", "20.000", 1), ("bester handyhülle pixel 8", "300", 1),
                      ("kopfhörer bluetooth", "30.000", 1), ("pixel 8 handyhülle schwarz matt", "1.200", 1)]:
                w.writerow(r)
        out = self.p("seo.json")
        run(os.path.join(PI, "seo_import.py"), de, "--marketplace", "DE", "--product-terms", "pixel 8,hülle,handyhülle",
            "--competitors", "samsung", "--seo-date", "2026-10-01", "--out", out)
        d = json.load(open(out, encoding="utf-8"))
        by = {k["keyword"]: k for k in d["keywords"]}
        self.assertEqual(by["handyhülle pixel 8"]["tier"], "TIER_1_PRIMARY")
        self.assertEqual(by["pixel 8 handyhülle schwarz matt"]["tier"], "TIER_3_LONG_TAIL")
        self.assertEqual(by["samsung galaxy hülle"]["status"], "COMPETITOR_TERM")
        self.assertEqual(by["bester handyhülle pixel 8"]["status"], "PROHIBITED_TERM")
        self.assertEqual(by["kopfhörer bluetooth"]["status"], "IRRELEVANT")
        self.assertEqual(sum(k["status"] == "SEMANTIC_DUPLICATE" for k in d["keywords"]), 1)
        self.assertEqual(d["sanitization_report"]["competitor_terms"], 1)
        # an English (US) export must never be accepted for DE
        us = self.p("cerebro_us.csv")
        with open(us, "w", newline="") as f:
            w = _csv.writer(f)
            w.writerow(["Keyword Phrase", "Search Volume"])
            for r in [("phone case for pixel 8", 12000), ("case with stand for women", 5000), ("clear case for men", 4000)]:
                w.writerow(r)
        p = run(os.path.join(PI, "seo_import.py"), us, "--marketplace", "DE", "--product-terms", "pixel 8", ok=(1,))
        self.assertIn("SEO_MARKETPLACE_MISMATCH", p.stdout)
        # without product terms nothing is tiered
        r = run(os.path.join(PI, "seo_import.py"), de, "--marketplace", "DE", "--seo-date", "2026-10-01", "--out", self.p("s2.json"))
        self.assertIn("NEEDS_PRODUCT_TERMS", json.dumps(json.load(open(self.p("s2.json"), encoding="utf-8"))))

    def test_countries_and_languages(self):
        import openpyxl as _x
        def mk(name, rows):
            wb = _x.Workbook(); ws = wb.active
            ws.append(["Helium 10 Cerebro"]); ws.append([]); ws.append(["Keyword Phrase", "Search Volume"])  # title rows above the header
            for r in rows:
                ws.append(r)
            wb.save(self.p(name)); return self.p(name)
        seo = os.path.join(PI, "seo_import.py")
        # multi-language marketplace: language must be chosen
        ca = mk("ca.xlsx", [("phone case for pixel 8", 9000), ("pixel 8 case with stand", 4000)])
        p = run(seo, ca, "--marketplace", "CA", "--product-terms", "pixel 8", ok=(1,))
        self.assertIn("ask the user which language", p.stderr + p.stdout)
        p = run(seo, ca, "--marketplace", "CA", "--language", "fr", "--product-terms", "pixel 8", ok=(1,))
        self.assertIn("SEO_LANGUAGE_MISMATCH", p.stdout)
        run(seo, ca, "--marketplace", "CA", "--language", "en", "--product-terms", "pixel 8", "--seo-date", "2026-10-01")
        # script-based languages
        jp = mk("jp.xlsx", [("Pixel 8 ケース", 15000), ("スマホケース 耐衝撃", 20000), ("phone case for pixel 8", 3000)])
        out = self.p("jp.json")
        run(seo, jp, "--marketplace", "JP", "--product-terms", "pixel 8,ケース", "--seo-date", "2026-10-01", "--out", out)
        by = {k["keyword"]: k for k in json.load(open(out, encoding="utf-8"))["keywords"]}
        self.assertEqual(by["Pixel 8 ケース"]["language"], "ja")
        self.assertEqual(by["phone case for pixel 8"]["status"], "WRONG_LANGUAGE")
        ae = mk("ae.xlsx", [("جراب بكسل 8", 3000), ("غطاء هاتف", 2000)])
        run(seo, ae, "--marketplace", "AE", "--language", "ar", "--product-terms", "بكسل,جراب,غطاء", "--seo-date", "2026-10-01")
        # locale lists for content checks
        p = run(os.path.join(PI, "content_check.py"), "--title", "Acme etui na telefon", "--lang", "pl")
        self.assertIn("LOCALE_LIST_MISSING", p.stdout)
        # handoff: language required for multi-language marketplaces, currency must match the marketplace
        rec = json.loads(run(os.path.join(PI, "handoff_tool.py"), "--example").stdout)
        rec["marketplace"] = "CA"
        rec["content"].pop("language")
        self.write_json("h_ca.json", rec)
        run(os.path.join(PI, "handoff_tool.py"), "seal", self.p("h_ca.json"), self.p("h_ca.jsonl"))
        p = run(os.path.join(PI, "handoff_tool.py"), "validate", self.p("h_ca.jsonl"), ok=(1,))
        self.assertIn("content.language is required for CA", p.stdout)
        self.assertIn("CURRENCY_CONFLICT", p.stdout)

    def test_marketplace_table_matches_pricing_engine(self):
        sys.path.insert(0, SH)
        import marketplaces, pricing_engine
        for code, cur in pricing_engine.MARKETPLACE_CURRENCY.items():
            self.assertEqual(marketplaces.currency(code), cur, code)
        self.assertEqual(marketplaces.languages("JP"), ["ja"])
        self.assertTrue(marketplaces.needs_language_choice("BE"))

    def test_google_link_cli(self):
        import http.server, io, threading, zipfile as _zf
        import openpyxl as _x

        def xl():
            wb = _x.Workbook(); wb.active.append(["Keyword Phrase", "Search Volume"]); b = io.BytesIO(); wb.save(b); return b.getvalue()

        class H(http.server.BaseHTTPRequestHandler):
            def log_message(self, *a): pass
            def reply(self, code, body=b"", ctype="text/plain", extra=None):
                self.send_response(code); self.send_header("Content-Type", ctype)
                for k, v in (extra or {}).items(): self.send_header(k, v)
                self.end_headers(); self.wfile.write(body)
            def do_GET(self):
                p = self.path
                if "PRIVATEFILE" in p and self.headers.get("Authorization") != "Bearer tok123": return self.reply(403, b"no")
                if "LOGINFILE" in p: return self.reply(302, extra={"Location": "https://accounts.google.com/ServiceLogin"})
                if "EVILFILE" in p: return self.reply(302, extra={"Location": "https://evil.example.com/x"})
                if "NOTNATIVE" in p and "/export" in p: return self.reply(400, b"bad")
                if "/export?format=xlsx" in p: return self.reply(200, xl(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                if "/export?format=csv" in p: return self.reply(200, ("gid=" + p.split("gid=")[-1] + "\na,b\n").encode(), "text/csv")
                if p.startswith("/uc?"):
                    z = io.BytesIO()
                    with _zf.ZipFile(z, "w") as zf: zf.writestr("xl/vbaProject.bin", b"x")
                    return self.reply(200, z.getvalue(), "application/octet-stream", {"Content-Disposition": 'attachment; filename="FeedTemplate.xlsm"'})
                self.reply(404, b"nf")

        srv = http.server.HTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        env = dict(os.environ, GOOGLE_LINK_TEST_BASE=f"http://127.0.0.1:{srv.server_address[1]}")
        gl = os.path.join(SH, "google_link.py")
        out = self.p("gdl")
        sid = "1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789"

        def go(url, *extra, token=None, ok=(0,)):
            e = dict(env, **({"GOOGLE_ACCESS_TOKEN": token} if token else {}))
            p = subprocess.run([sys.executable, gl, "fetch", url, "--out", out, "--json", *extra], capture_output=True, text=True, env=e)
            self.assertIn(p.returncode, ok, p.stdout + p.stderr)
            return p
        info = json.loads(run(gl, "info", f"https://docs.google.com/spreadsheets/d/{sid}/edit#gid=77").stdout)
        self.assertEqual((info["kind"], info["gid"]), ("sheet", "77"))
        r = json.loads(go(f"https://docs.google.com/spreadsheets/d/{sid}/edit").stdout)
        self.assertTrue(r["path"].endswith(".xlsx") and r["warnings"])  # warns: not an Amazon feed template
        r = json.loads(go(f"https://docs.google.com/spreadsheets/d/{sid}/edit#gid=77", "--format", "csv").stdout)
        self.assertIn("gid=77", open(r["path"], encoding="utf-8").read())
        r = json.loads(go(f"https://drive.google.com/file/d/{sid}/view").stdout)
        self.assertTrue(r["path"].endswith("FeedTemplate.xlsm"))  # original file keeps its name and extension
        self.assertIn("403", go("https://docs.google.com/spreadsheets/d/PRIVATEFILE0123456789abcdef/edit", ok=(1,)).stderr)
        go("https://docs.google.com/spreadsheets/d/PRIVATEFILE0123456789abcdef/edit", token="tok123")
        self.assertIn("NOT_SHARED", go("https://docs.google.com/document/d/LOGINFILE0123456789abcdef/edit", ok=(1,)).stderr)
        self.assertIn("non-Google host refused", go("https://docs.google.com/document/d/EVILFILE0123456789abcdef/edit", ok=(1,)).stderr)
        r = json.loads(go("https://docs.google.com/spreadsheets/d/NOTNATIVE0123456789abcdef/edit").stdout)
        self.assertEqual(r["format"], "raw")  # an uploaded Office file opened in Sheets: falls back to the Drive download
        for bad in ("https://evil.example.com/spreadsheets/d/" + sid, "http://docs.google.com/document/d/" + sid,
                    "https://drive.google.com/drive/folders/" + sid):
            self.assertEqual(go(bad, ok=(1,)).returncode, 1)
        srv.shutdown()

    def test_amazon_link_cli(self):
        al = os.path.join(SH, "amazon_link.py")
        def parse(url, ok=(0,)):
            p = run(al, "parse", url, ok=ok)
            return json.loads(p.stdout) if p.returncode == 0 else p
        d = parse("https://www.amazon.de/-/en/Name/dp/b0chx3qbch/ref=sr_1_3?qid=1&sr=8-3&tag=aff-21&th=1")
        self.assertEqual((d["marketplace"], d["kind"], d["asin"], d["language_hint"]), ("DE", "product", "B0CHX3QBCH", "en"))
        self.assertEqual(d["canonical_url"], "https://www.amazon.de/dp/B0CHX3QBCH")
        self.assertIn("tag", d["tracking_params_removed"])
        self.assertTrue(d["note_affiliate"])
        d = parse("https://www.amazon.co.uk/product-reviews/B0CHX3QBCH/?ie=UTF8")
        self.assertEqual((d["marketplace"], d["subkind"]), ("UK", "reviews"))
        self.assertTrue(d["canonical_url"].endswith("/product-reviews/B0CHX3QBCH"))
        d = parse("https://www.amazon.com/Cases/b/ref=x?ie=UTF8&node=2407760011")
        self.assertEqual((d["kind"], d["browse_nodes"], d["canonical_url"]), ("category", ["2407760011"], "https://www.amazon.com/b?node=2407760011"))
        d = parse("https://www.amazon.com/Best-Sellers/zgbs/wireless/2407760011/ref=zg")
        self.assertEqual((d["kind"], d["board"], d["browse_nodes"]), ("ranking", "bestsellers", ["2407760011"]))
        d = parse("https://www.amazon.co.jp/s?k=pixel+8+%E3%82%B1%E3%83%BC%E3%82%B9&rh=n%3A2285178051")
        self.assertEqual((d["marketplace"], d["keyword"], d["browse_nodes"]), ("JP", "pixel 8 ケース", ["2285178051"]))
        for host, code in (("amazon.pl", "PL"), ("amazon.se", "SE"), ("amazon.com.tr", "TR"), ("amazon.ae", "AE"), ("amazon.in", "IN"),
                           ("amazon.com.be", "BE"), ("amazon.sa", "SA"), ("amazon.com.au", "AU"), ("amazon.sg", "SG")):
            self.assertEqual(parse(f"https://www.{host}/dp/B0CHX3QBCH")["marketplace"], code)
        for bad in ("https://www.amazon.de.evil.com/dp/B0CHX3QBCH", "https://evil.com/dp/B0CHX3QBCH", "https://www.amazon.de/gp/cart/view.html",
                    "https://www.amazon.de/ap/signin", "https://www.amazon.xx/dp/B0CHX3QBCH"):
            self.assertIn("REFUSED", parse(bad, ok=(1,)).stderr)
        self.assertEqual(run(al, "build", "--marketplace", "JP", "--keyword", "pixel 8 ケース").stdout.strip(),
                         "https://www.amazon.co.jp/s?k=pixel+8+%E3%82%B1%E3%83%BC%E3%82%B9")
        rows = run(al, "all-markets", "--asin", "B0CHX3QBCH", "--marketplaces", "DE,CA,AE").stdout.strip().splitlines()
        self.assertEqual([r.split("\t")[0] for r in rows], ["DE", "CA", "AE"])
        plan = json.loads(run(al, "plan", "https://www.amazon.de/dp/B0CHX3QBCH").stdout)
        tools = {s["tool"].split("__")[-1]: s for s in plan["data_sources"]}
        self.assertTrue(tools["get_listing_details"]["marketplace_supported"])
        self.assertEqual(tools["get_keywords_by_asin"]["arguments"]["asin"], "B0CHX3QBCH")
        se = json.loads(run(al, "plan", "https://www.amazon.se/dp/B0CHX3QBCH").stdout)
        self.assertFalse(se["data_sources"][0]["marketplace_supported"])  # SE is not offered by the MCP tools: say so, never substitute
        cat = json.loads(run(al, "plan", "https://www.amazon.com/b?node=2407760011").stdout)
        self.assertEqual(cat["data_sources"][0]["arguments"]["filters"]["category"], ["2407760011"])
        self.assertEqual(run(al, "open", "--marketplace", "FR", "--asin", "B0CHX3QBCH", "--print-only").stdout.strip(), "https://www.amazon.fr/dp/B0CHX3QBCH")

    def test_competitor_reference_pack_and_copy_guard(self):
        import http.server, io as _io, threading
        from PIL import Image as _I

        def png(c):
            b = _io.BytesIO(); _I.new("RGB", (400, 400), c).save(b, "PNG"); return b.getvalue()

        class H(http.server.BaseHTTPRequestHandler):
            def log_message(self, *a): pass
            def do_GET(self):
                ok = self.path.startswith("/img/")
                self.send_response(200); self.send_header("Content-Type", "image/png" if ok else "text/html"); self.end_headers()
                self.wfile.write(png("white" if "main" in self.path else (30, 90, 160)) if ok else b"<html>")

        srv = http.server.HTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        port = srv.server_address[1]
        listing = {"data": {"images": [{"url": f"http://127.0.0.1:{port}/img/main.png", "variant": "MAIN"},
                                       {"url": f"http://127.0.0.1:{port}/img/pt01.png", "variant": "PT01"},
                                       {"url": "https://evil.example.com/x.png", "variant": "PT02"},
                                       {"url": f"http://127.0.0.1:{port}/page.html", "variant": "PT03"}],
                            "product_name": "BrandX Slim Case for Pixel 8, Shockproof Matte Black Cover, 6.2 inch, MIL-STD-810G",
                            "bullet_points": ["MILITARY GRADE PROTECTION: tested to MIL-STD-810G drop standards, keeps your phone safe from falls up to 3 m",
                                              "Precise cutouts for buttons and ports"],
                            "description": "BrandX slim case offers waterproof style protection"}}
        lj = self.write_json("cp_listing.json", listing)
        dj = self.write_json("cp_details.json", {"data": {"price": 2499, "best_sellers_rank": 1520, "review_count": 4312}})
        cp = os.path.join(PI, "competitor_pack.py")
        out = self.p("cp_out")
        env = dict(os.environ, COMPETITOR_PACK_TEST_ALLOW_HOST="127.0.0.1")
        p = subprocess.run([sys.executable, cp, "ingest", lj, dj, "--marketplace", "DE", "--asin", "B0CHX3QBCH", "--out", out,
                            "--download-images", "--delay", "0"], capture_output=True, text=True, env=env)
        self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
        card_path = os.path.join(out, "competitor_B0CHX3QBCH_DE.json")
        card = json.load(open(card_path, encoding="utf-8"))
        self.assertEqual(card["use"], "REFERENCE_ONLY")
        self.assertIn("IP_REVIEW_REQUIRED", card["flags"])
        self.assertEqual(card["observations"]["brand_guess"], "BrandX")
        self.assertIn("waterproof", card["observations"]["claims_needing_evidence"])
        self.assertIn("MIL-STD-810G", card["observations"]["certification_mentions"])
        self.assertNotIn("810 g", card["observations"]["measurable_features"])  # part of a standard name, not a weight
        by = {i["variant"]: i for i in card["images"]}
        self.assertTrue(os.path.exists(by["MAIN"]["path"]) and by["MAIN"]["white_background_corners"])
        self.assertIn("host not an Amazon image host", by["PT02"]["skipped"])  # never fetch from other hosts
        self.assertIn("not an image", by["PT03"]["error"])
        self.assertEqual(card["metrics"]["bsr"], 1520)
        # copy guard
        copied = self.write_json("cp_copied.json", {"title": "x", "bullet_points": [listing["data"]["bullet_points"][0]]})
        own = self.write_json("cp_own.json", {"title": "Acme FlexShield cover", "bullet_points": ["Corner airbags absorb impact when dropped from table height"]})
        run(cp, "similarity", card_path, copied, ok=(1,))
        run(cp, "similarity", card_path, own)
        self.assertEqual(run(cp, "brands", card_path).stdout.strip(), "BrandX")
        srv.shutdown()

    def test_content_check(self):
        p = run(os.path.join(PI, "content_check.py"), "--title",
                "Acme Case Case Case for Phone best price 9,99 EUR", ok=(1,))
        self.assertIn("TITLE_WORD_REPETITION", p.stdout)
        self.assertIn("FORBIDDEN_TERM_FOUND", p.stdout)

    def test_category_and_research_references_shipped(self):
        for a in ("amazon-product-intelligence", "amazon-feed-compiler", "amazon-feed-error-agent"):
            refs = os.path.join(ROOT, "plugins", a, "skills", a, "references")
            for f in ("amazon-categories.md", "live-research.md"):
                self.assertGreater(os.path.getsize(os.path.join(refs, f)), 1000, (a, f))
            skill = open(os.path.join(refs, "..", "SKILL.md"), encoding="utf-8").read()
            self.assertIn("amazon-categories", skill)
            self.assertIn("live-research", skill)

    def test_content_limits_2026(self):
        cc = os.path.join(PI, "content_check.py")
        title75 = "A" * 75
        good = run(cc, "--title", "Acme Phone Case iPhone 15 Pro", "--highlights", "Shockproof, clear, MagSafe ready",
                   "--backend", "huelle schutz", "--marketplace", "DE", "--brand", "Acme", "--json")
        self.assertNotIn("TITLE_TOO_LONG", good.stdout)
        p = run(cc, "--title", "x" * 76, ok=(1,))
        self.assertIn("TITLE_TOO_LONG", p.stdout)
        run(cc, "--title", "x" * 76, "--category", "media", ok=(0,))
        # backend is measured in bytes and differs per marketplace
        cyr = "кейс " * 30  # 150 chars, 270 bytes
        self.assertIn("BACKEND_TOO_LONG", run(cc, "--title", "Acme", "--backend", cyr, "--marketplace", "DE", ok=(1,)).stdout)
        self.assertNotIn("BACKEND_TOO_LONG", run(cc, "--title", "Acme", "--backend", cyr, "--marketplace", "JP", ok=(0, 1)).stdout)
        self.assertIn("BACKEND_TOO_LONG", run(cc, "--title", "Acme", "--backend", "a" * 210, "--marketplace", "IN", ok=(1,)).stdout)
        self.assertIn("BACKEND_COMMAS", run(cc, "--title", "Acme", "--backend", "case, cover", ok=(1,)).stdout)
        # item highlights must not repeat title words
        p = run(cc, "--title", "Acme Magnetic Phone Case", "--highlights", "Magnetic phone case, clear cover", ok=(1,))
        self.assertIn("HIGHLIGHTS_REPEAT_TITLE", p.stdout)
        # bullets and description
        rec = self.write_json("cl_rec.json", {"title": "Acme Case", "bullet_points": ["Short", "ok bullet text here.", "\u2705 check mark bullet"],
                                              "description": "<p>html</p>" + "d" * 2000})
        p = run(cc, "--file", rec, ok=(1,))
        for code in ("BULLET_TOO_SHORT", "BULLET_END_PUNCTUATION", "BULLET_SYMBOLS", "DESCRIPTION_TOO_LONG", "DESCRIPTION_HTML"):
            self.assertIn(code, p.stdout)
        run(cc, "--title", "Acme", "--bullets-max", "5", "--marketplace", "DE")  # no crash without highlights/bullets

    def test_marketplace_languages_primary_secondary(self):
        if SH not in sys.path:
            sys.path.insert(0, SH)
        import marketplaces
        self.assertEqual(marketplaces.languages("EG"), ["ar"])
        self.assertEqual(marketplaces.secondary_languages("DE"), ["en", "nl", "pl", "tr"])
        self.assertTrue(marketplaces.needs_language_choice("AE"))
        self.assertFalse(marketplaces.needs_language_choice("IN"))
        self.assertFalse(marketplaces.needs_language_choice("US"))

    def test_seo_export_marketplace_is_asked_not_guessed(self):
        seo = os.path.join(PI, "seo_import.py")
        kw = os.path.join(self.d, "iphone_export.csv")
        with open(kw, "w", encoding="utf-8") as fh:
            fh.write("Keyword Phrase,Search Volume,Keyword Sales\niphone 17 case,100000,500\niphone 17 pro case,80000,400\n")
        p = run(seo, kw, "--marketplace", "ES", "--product-terms", "iphone", ok=(1,))
        self.assertIn("SEO_SOURCE_MARKETPLACE_UNKNOWN", p.stdout)
        p = run(seo, kw, "--marketplace", "ES", "--product-terms", "iphone", "--source-marketplace", "US", ok=(1,))
        self.assertIn("SEO_MARKETPLACE_MISMATCH", p.stdout)
        run(seo, kw, "--marketplace", "US", "--product-terms", "iphone", "--source-marketplace", "US")
        named = os.path.join(self.d, "GB_body_lotion.csv")
        with open(named, "w", encoding="utf-8") as fh:
            fh.write("Keyword Phrase,Search Volume\nbody lotion,5000\nmoisturiser for dry skin,3000\n")
        run(seo, named, "--marketplace", "UK", "--product-terms", "lotion,moisturiser")  # GB in file name = UK
        p = run(seo, named, "--marketplace", "DE", "--product-terms", "lotion", ok=(1,))
        self.assertIn("SEO_MARKETPLACE_MISMATCH", p.stdout)

    def test_analytics_import_reports(self):
        an = os.path.join(PI, "analytics_import.py")
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Display Order", "Product Details", "ASIN", "URL", "Brand", "Price  \u20ac", "ASIN Sales", "ASIN Revenue", "Title Char. Count",
                   "BSR", "Ratings", "Review Count", "Images", "Fulfillment", "Creation Date", "Sponsored"])
        ws.append(["1.", "Alpha Dog Oil", "B000000001", "https://www.amazon.de/dp/B000000001", "Alpha", "16,99", "45", "1\u00a0015,50", "90",
                   "13\u00a0666", "4,2", "264", "7", "FBA", "Jan 4, 2022", ""])
        ws.append(["2.", "Beta Dog Oil", "B000000002", "https://www.amazon.de/dp/B000000002", "Beta", "24,90", "10", "249", "60",
                   "5\u00a0000", "4,5", "19", "6", "FBM", "Jan 4, 2022", ""])
        x = self.p("xray.xlsx")
        wb.save(x)
        p = run(an, x, "--out", self.p("an.json"), ok=(0, 1))
        rep = json.load(open(self.p("an.json"), encoding="utf-8"))["reports"][0]
        self.assertEqual((rep["kind"], rep["marketplace"], rep["currency"]), ("XRAY", "DE", "\u20ac"))
        self.assertEqual(rep["price"]["max"], 24.9)
        self.assertEqual(rep["monthly_revenue_total"], 1264.5)  # "1 015,50" read with the German format
        self.assertEqual(rep["title_over_75"], 1)
        self.assertIn("ESTIMATES", rep["source"])
        self.assertEqual(rep["category_scope"], "UNKNOWN")
        run(an, x, "--out", self.p("an_c.json"), "--category", "Pet Supplies > Dogs > Vitamins", ok=(0, 1))
        self.assertIn("LEAF", json.load(open(self.p("an_c.json"), encoding="utf-8"))["reports"][0]["category_scope"])
        cat = self.p("GB_catalog.csv")
        with open(cat, "w", encoding="utf-8-sig") as fh:
            fh.write('Reporting range=["Quarterly"],Select year=["2026"],Select quarter=["1"],Brand=["Acme"],Category=["Beauty"]\n'
                     '"ASIN title",ASIN,Category,"Impressions: Impressions","Clicks: Clicks","Basket Adds: Basket Adds","Purchases: Purchases","Reporting Date"\n'
                     '"A","B0AAAAAAA1","Beauty","1000","50","5","0","2026-03-31"\n"B","B0AAAAAAA2","Beauty","1000","50","5","10","2026-03-31"\n')
        sqp = self.p("GB_sqp.csv")
        with open(sqp, "w", encoding="utf-8-sig") as fh:
            fh.write('Brand=["Acme"],Reporting range=["Quarterly"],Select year=["2025"],Select quarter=["1"]\n'
                     '"Search query","Search query score","Search query volume","Impressions: Brand Share %","Impressions: Brand Count","Clicks: Brand Count","Purchases: Brand Count"\n'
                     '"camel milk cream","3","900","1.5","10","4","0"\n')
        seo = self.p("sqp_seo.csv")
        p = run(an, cat, sqp, "--out", self.p("an2.json"), "--emit-seo-csv", seo, ok=(1,))
        self.assertIn("PERIOD_MISMATCH", p.stdout)
        data = json.load(open(self.p("an2.json"), encoding="utf-8"))
        cat_items = {i["asin"]: i for i in next(r for r in data["reports"] if r["kind"] == "BA_CATALOG")["items"]}
        self.assertIn("CLICKS_NO_PURCHASE", " ".join(cat_items["B0AAAAAAA1"]["diagnosis"]))
        self.assertEqual(data["reports"][0]["marketplace"], "UK")  # GB_ in the file name
        self.assertIn("camel milk cream", open(seo, encoding="utf-8").read())
        run(os.path.join(PI, "seo_import.py"), seo, "--marketplace", "UK", "--source-marketplace", "UK", "--product-terms", "cream")
        unk = self.p("whatever.csv")
        open(unk, "w").write("a,b\n1,2\n")
        self.assertIn("UNKNOWN_REPORT", run(an, unk, ok=(1,)).stdout)
        nomp = self.p("catalog_nomp.csv")
        open(nomp, "w", encoding="utf-8").write(open(cat, encoding="utf-8-sig").read())
        self.assertIn("MARKETPLACE_UNKNOWN", run(an, nomp, ok=(1,)).stdout)

    def test_processing_summary_status_layout_and_change_plan(self):
        """Current Seller Central report: 3 leading status columns, per-code and per-SKU tables, impacted cell = status cell."""
        import openpyxl
        keys = ["contribution_sku#1.value", "product_type#1.value", "generic_keyword[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
                "other_product_image_locator_1[marketplace_id=ATVPDKIKX0DER]#1.media_location"]

        def template(with_status):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Template"
            pre = ["::submission_status", "::number_of_attributes_with_errors", "::number_of_attributes_with_other_suggestions"] if with_status else []
            ws.append(["settings=feedType=256&timestamp=2026-10-11T10%3A00%3A00.000Z"])
            ws.append(["Use ENGLISH"])
            ws.append([])
            ws.append(["Status"] * len(pre) + ["SKU", "Product Type", "Keywords", "Image 1"])
            ws.append(pre + keys)
            ws.append(([None] * len(pre)) + ["ABC123", "ACCESSORY", "", ""])
            rows = [("S1", "air pod procase pods", "https://h.example/a.jpg"), ("S2", "case cover", "https://h.example/b.jpg")]
            for sku, kw, url in rows:
                st = (["SUCCESS (OTHER ERROR)", 1, 0] if sku == "S1" else ["SUCCESS", 0, 0]) if with_status else []
                ws.append(st + [sku, "PORTABLE_ELECTRONIC_DEVICE_COVER", kw, url])
            return wb
        rep, feed = self.p("rep_status.xlsx"), self.p("orig_feed.xlsx")
        wb = template(True)
        sm = wb.create_sheet("Feed Processing Summary", 0)
        for r in ([], [], ["Summary"], ["Number of SKUs processed", "", "", 2], ["Number of SKUs successful", "", "", 1],
                  ["Number of SKUs successful with other errors", "", "", 1], [],
                  ["Errors and Warnings per Error Code"],
                  ["#", "Error code", "Category of error", "Store", "Error message", "Affected field (Impacted column)", "Number of errors"],
                  [1, "5561", "ERROR", "US", "Your generic_keyword contains Trademarked terms [PROCASE]. Remove any trademark terms and resubmit your listing.",
                   "Number of attributes with errors (B)", 1],
                  [2, "20000", "ERROR", "US", "Your media at https://h.example/a.jpg couldn't be downloaded due to a timeout.", "Number of attributes with errors (B)", 1],
                  ["Total number of errors", 2], [],
                  ["Errors and Warnings per SKU"],
                  ["#", "Error code", "Category of error", "Store", "Error message", "Affected field (Impacted cell)", "SKU"],
                  [1, "5561", "ERROR", "US", "Your generic_keyword contains Trademarked terms [PROCASE]. Remove any trademark terms and resubmit your listing.",
                   "Number of attributes with errors (B7)", "S1"],
                  [2, "20000", "ERROR", "US", "Your media at https://h.example/a.jpg couldn't be downloaded due to a timeout.",
                   "Number of attributes with errors (B7)", "S1"]):
            sm.append(r)
        wb.save(rep)
        template(False).save(feed)
        parsed = self.p("status.json")
        p = run(os.path.join(ER, "parse_processing_report.py"), rep, "--marketplace", "US", "--json", parsed)
        self.assertIn("STATUS layout", p.stdout)
        d = json.load(open(parsed, encoding="utf-8"))
        by = {f["error_code"]: f for f in d["findings"]}
        self.assertEqual(by["5561"]["proposed_value"], "air pod pods")
        self.assertTrue(by["5561"]["attribute"].startswith("generic_keyword"))
        self.assertTrue(by["20000"]["attribute"].startswith("other_product_image_locator_1"))
        self.assertEqual(d["submission_status_counts"], {"SUCCESS (OTHER ERROR)": 1, "SUCCESS": 1})
        self.assertEqual(by["5561"]["cell"], "F7")  # report letters include the 3 status columns ...
        plan = self.p("plan.json")
        cells = self.p("plan_cells.json")
        run(os.path.join(ER, "build_change_plan.py"), parsed, "--feed", feed, "--out", plan, "--approve", "ALL", "--cells-out", cells)
        pl = json.load(open(plan, encoding="utf-8"))
        chg = [i for i in pl["items"] if i["kind"] == "CELL_CHANGE"]
        self.assertEqual((chg[0]["cell"], chg[0]["proposed"]), ("C7", "air pod pods"))  # ... the original feed does not
        self.assertEqual([i["error_code"] for i in pl["items"] if i["kind"] == "ACTION_ONLY"], ["20000"])
        out = self.p("orig_feed_v2.xlsx")
        run(os.path.join(ER, "xlsm_patch.py"), "--source", feed, "--cells", cells, "--out", out, "--sheet", "Template", "--min-row", "7")
        run(os.path.join(ER, "workbook_guard.py"), "--source", feed, "--output", out, "--sheet", "Template", "--approved", cells)
        self.assertEqual(openpyxl.load_workbook(out)["Template"]["C7"].value, "air pod pods")
        # same upload twice is not a re-upload
        self.assertIn("SAME_REPORT", run(os.path.join(ER, "compare_reports.py"), parsed, parsed, ok=(1,)).stdout)
        # a changed feed (value differs from the report) is not patched
        wb2 = openpyxl.load_workbook(feed)
        wb2["Template"]["C7"] = "edited later"
        edited = self.p("edited.xlsx")
        wb2.save(edited)
        p = run(os.path.join(ER, "build_change_plan.py"), parsed, "--feed", edited, "--out", self.p("plan2.json"), ok=(1,))
        self.assertIn("'mismatches': 1", p.stdout)


if __name__ == "__main__":
    unittest.main()
