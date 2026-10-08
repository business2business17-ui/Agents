#!/usr/bin/env python3
"""Build the separate XLSX error report + Change Plan sheet (spec sections 8, 9, 40).

Usage:
  build_error_report.py parsed.json report.xlsx [--plan change_plan.json] [--history history.json]
parsed.json  = output of parse_processing_report.py
change_plan  = [ {"change_id":"CHG-001","cell":"E7","attribute":"color_name","current":"Bluee","proposed":"Blue",
                  "error_code":"8058","reason":"...","source":"Valid Values","confidence":"HIGH",
                  "root_cause":"INVALID_ENUM","user_decision":"PENDING","protected":false} ]
Sheets: ERRORS (spec columns, statuses NEW/ANALYZED/WAITING_USER/APPROVED/PATCHED/UPLOADED/RESOLVED/REJECTED_AGAIN/
ESCALATION_REQUIRED), CHANGE_PLAN (Before/After with confidence and approval state), SUMMARY.
Known error codes are enriched from assets/error_kb.json. Nothing here edits a feed.
"""
import argparse
import json
import os
import sys
from collections import Counter

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

KB = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "error_kb.json")
COLS = ["Marketplace", "Error Fingerprint", "Error Code", "Error Category", "Severity", "Business Priority", "SKU",
        "EAN / UPC / GTIN", "ASIN", "Template Sheet", "Template Row", "Cell", "Column", "Attribute", "Original Value",
        "Amazon Message", "Root Cause", "Proposed Value", "Solution", "Source", "Confidence", "User Decision", "Applied",
        "Attempt", "Result After Upload", "Status"]


def head(ws, header):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(12, len(h) + 2))
    ws.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("parsed")
    ap.add_argument("out")
    ap.add_argument("--plan")
    ap.add_argument("--history")
    a = ap.parse_args()
    parsed = json.load(open(a.parsed, encoding="utf-8"))
    kb = json.load(open(KB, encoding="utf-8"))["codes"]
    plan = json.load(open(a.plan, encoding="utf-8")) if a.plan else []
    hist = json.load(open(a.history, encoding="utf-8")) if a.history else {}
    plan_by_cell = {p["cell"]: p for p in plan}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ERRORS"
    head(ws, COLS)
    st = Counter()
    for f in parsed["findings"]:
        code = str(f.get("error_code") or "")
        e = kb.get(code, {})
        p = plan_by_cell.get(f.get("cell") or "", {})
        k = "|".join(str(f.get(x)) for x in ("marketplace", "sku", "attribute"))
        attempts = hist.get(k, 1)
        if attempts >= 3:
            status = "ESCALATION_REQUIRED"
        elif p.get("user_decision") == "APPROVED":
            status = "APPROVED"
        elif p:
            status = "WAITING_USER"
        else:
            status = "ANALYZED" if e else "NEW"
        st[status] += 1
        ws.append([f.get("marketplace"), f.get("fingerprint"), code, f.get("category") or e.get("type"), f.get("severity"),
                   "HIGH" if f.get("severity") == "ERROR" else "MEDIUM", f.get("sku"), ", ".join(f.get("ids") or []), None,
                   parsed.get("template_sheet"), f.get("row"), f.get("cell"), (f.get("cell") or "").rstrip("0123456789"),
                   f.get("attribute"), f.get("original_value"), f.get("error_message"), p.get("root_cause") or e.get("type"),
                   p.get("proposed"), p.get("reason") or "; ".join(e.get("solution", [])), p.get("source"), p.get("confidence"),
                   p.get("user_decision"), None, attempts, None, status])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(c.value, str):
                c.data_type = "s"
    pl = wb.create_sheet("CHANGE_PLAN")
    head(pl, ["Change ID", "SKU/EAN", "Sheet", "Cell", "Attribute", "Current Value", "Error", "Proposed Value", "Reason",
              "Confidence", "Protected field", "Action / User decision"])
    for p in plan:
        pl.append([p.get("change_id"), p.get("sku"), parsed.get("template_sheet"), p.get("cell"), p.get("attribute"), p.get("current"),
                   p.get("error_code"), p.get("proposed"), p.get("reason"), p.get("confidence"),
                   "YES - separate approval" if p.get("protected") else "no", p.get("user_decision", "PENDING")])
    sm = wb.create_sheet("SUMMARY")
    head(sm, ["Metric", "Value"])
    t = parsed.get("totals", {})
    for k, v in [("Errors found", t.get("errors")), ("Warnings", t.get("warnings")), ("Unmapped summary lines", t.get("unknown")),
                 ("Proposed changes", len(plan)),
                 ("HIGH confidence", sum(p.get("confidence") == "HIGH" for p in plan)),
                 ("MEDIUM confidence", sum(p.get("confidence") == "MEDIUM" for p in plan)),
                 ("LOW confidence", sum(p.get("confidence") == "LOW" for p in plan)),
                 ("Protected-field changes", sum(bool(p.get("protected")) for p in plan)),
                 ("Escalation required", st["ESCALATION_REQUIRED"])]:
        sm.append([k, v])
    wb.save(a.out)
    print(f"wrote {a.out}: {len(parsed['findings'])} findings, {len(plan)} planned changes; status counts {dict(st)}")


if __name__ == "__main__":
    main()
