#!/usr/bin/env python3
"""Match product images to XLSX rows by GTIN / EAN / UPC (image file name = code).

Usage:
  match_inputs.py --images DIR --xlsx FILE [--sheet NAME] [--recursive] [--out match.json]

Image names: `4006381333931.jpg`, `036000291452.png`, and extra angles of the same product
`4006381333931_2.jpg`, `4006381333931-back.png`, `4006381333931 (1).jpg` (suffix = view label).
XLSX: one row per product; the GTIN column is found by header (gtin/ean/upc/barcode/штрихкод...) or,
failing that, by content (column of 8-14 digit codes). Columns whose header means "benefits/advantages"
are separated from the technical facts (TTX). Everything else is kept as facts under its own header.

Matching key = digits zero-padded to 14 (so UPC-12 == EAN-13 with a leading 0, and a code that lost its
leading zero in Excel still matches if the padded value has a valid check digit; flagged ZERO_PADDED).
Nothing is guessed: unmatched images, unmatched rows, duplicates and invalid codes are reported, never fixed.

Exit code 0 = everything matched cleanly, 1 = something needs review (see JSON `issues`), 2 = cannot run.
The source XLSX and images are only read.
"""
import argparse
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")
try:
    from PIL import Image
except ImportError:  # pixel size is a nice-to-have
    Image = None

IMG_EXT = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp", ".bmp"}
ID_HEADER = re.compile(r"gtin|\bean\b|\bupc\b|barcode|bar code|штрих|баркод|штрих-код|\bean13\b|external.?product.?id", re.I)
BENEFIT_HEADER = re.compile(
    r"преимущ|выгод|benefit|advantage|vorteil|avantage|ventaja|vantagg|selling.?point|\busp\b|key.?feature.?benefit",
    re.I)
CATEGORY_HEADER = re.compile(r"категор|category|kategorie|catégorie|categoría|categoria|department|раздел", re.I)
TYPE_HEADER = re.compile(r"тип\b|вид\b|тип товара|product.?type|item.?type|typ\b|type\b|tipo|sub.?categor|подкатегор", re.I)
DESC_HEADER = re.compile(r"описани|description|beschreibung|descripci|descrizione|аннотац|bullet|буллет|title|название|наименован|заголовок", re.I)
EMPTY_BENEFIT = {"", "-", "—", "–", "n/a", "na", "none", "нет", "нету", "null", "tbd", "?", "отсутствует"}
NAME_SPLIT = re.compile(r"^(\d{6,14})(?:[\s_\-.()]+(.*))?$")


def check_ok(body_plus_check):
    s = body_plus_check
    total = sum(int(ch) * (3 if i % 2 == 0 else 1) for i, ch in enumerate(reversed(s[:-1])))
    return (10 - total % 10) % 10 == int(s[-1])


def gtin_info(digits):
    """-> (key14 or None, status, note). Never changes the code, only normalizes the comparison key."""
    if not digits or not digits.isdigit():
        return None, "GTIN_INVALID", "not a digit string"
    if len(digits) > 14:
        return None, "GTIN_INVALID", f"{len(digits)} digits; expected 8, 12, 13 or 14"
    length = next(n for n in (8, 12, 13, 14) if len(digits) <= n)  # smallest valid length that fits; no other guesses
    padded = digits.zfill(length)
    if check_ok(padded):
        note = None if len(digits) == length else f"{len(digits)} digits: leading zero lost, zero-padded to {length} (check digit valid)"
        return digits.zfill(14), "GTIN_VALID" if note is None else "ZERO_PADDED", note
    return digits.zfill(14), "GTIN_INVALID", f"check digit mismatch for {length}-digit code (should end in {_cd(padded[:-1])})"


def _cd(body):
    total = sum(int(ch) * (3 if i % 2 == 0 else 1) for i, ch in enumerate(reversed(body)))
    return (10 - total % 10) % 10


def scan_images(folder, recursive):
    it = folder.rglob("*") if recursive else folder.glob("*")
    out, bad = [], []
    for p in sorted(it):
        if not p.is_file() or p.suffix.lower() not in IMG_EXT:
            continue
        m = NAME_SPLIT.match(p.stem.strip())
        if not m:
            bad.append({"path": str(p), "issue": "FILENAME_NOT_A_GTIN", "detail": f"'{p.stem}' does not start with 6-14 digits"})
            continue
        digits, label = m.group(1), (m.group(2) or "").strip() or None
        key, status, note = gtin_info(digits)
        rec = {"path": str(p), "gtin_raw": digits, "view_label": label, "key": key, "gtin_status": status, "note": note}
        if Image is not None:
            try:
                with Image.open(p) as im:
                    rec.update(width=im.width, height=im.height, mode=im.mode)
            except Exception as e:  # noqa: BLE001 - report, don't die
                rec["note"] = ((note + "; ") if note else "") + f"cannot open image: {e}"
        out.append(rec)
    return out, bad


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(v)
    if isinstance(v, int):
        return str(v)
    return str(v).strip()


def header_row_index(rows):
    best, best_n = 0, -1
    for i, r in enumerate(rows[:15]):
        n = sum(1 for c in r if isinstance(c, str) and c.strip())
        if n > best_n:
            best, best_n = i, n
    return best


def find_gtin_col(headers, body):
    for j, h in enumerate(headers):
        if h and ID_HEADER.search(h):
            return j, "header"
    best, best_ratio = None, 0.0
    for j in range(len(headers)):
        vals = [cell_text(r[j]) for r in body if j < len(r) and cell_text(r[j])]
        if len(vals) < 1:
            continue
        ok = sum(1 for v in vals if re.fullmatch(r"\d{8,14}", v.replace(" ", "")))
        ratio = ok / len(vals)
        if ratio > best_ratio:
            best, best_ratio = j, ratio
    return (best, "content") if best is not None and best_ratio >= 0.6 else (None, None)


def split_benefits(text):
    parts = re.split(r"[\n\r]+|\s*[•·▪●■]\s*|;\s+", text)
    return [re.sub(r"^[\-–—*\d.)\s]+", "", p).strip() for p in parts if p and p.strip()]


def read_rows(xlsx, sheet_name):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    names = [sheet_name] if sheet_name else wb.sheetnames
    found = []
    for name in names:
        if name not in wb.sheetnames:
            raise SystemExit(f"sheet '{name}' not found; sheets: {wb.sheetnames}")
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
        rows = [r for r in rows if any(c not in (None, "") for c in r)]
        if len(rows) < 2:
            continue
        hi = header_row_index(rows)
        headers = [cell_text(c) for c in rows[hi]]
        body = rows[hi + 1:]
        col, how = find_gtin_col(headers, body)
        if col is not None:
            found.append((name, hi, headers, body, col, how))
    return found


def main():
    ap = argparse.ArgumentParser(description="Match GTIN-named images to XLSX rows")
    ap.add_argument("--images", required=True, help="folder with images named <GTIN>[_suffix].ext")
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sheet")
    ap.add_argument("--recursive", action="store_true")
    ap.add_argument("--out")
    a = ap.parse_args()

    folder, xlsx = Path(a.images), Path(a.xlsx)
    if not folder.is_dir() or not xlsx.is_file():
        print("images folder or xlsx not found", file=sys.stderr)
        return 2
    images, bad_names = scan_images(folder, a.recursive)
    sheets = read_rows(xlsx, a.sheet)
    if not sheets:
        print("no sheet with a recognizable GTIN/EAN/UPC column and data rows", file=sys.stderr)
        return 2

    issues = [dict(b) for b in bad_names]
    rows_by_key, seen_rows = {}, []
    mapping = []
    for name, hi, headers, body, gcol, how in sheets:
        ben_cols = [j for j, h in enumerate(headers) if h and j != gcol and BENEFIT_HEADER.search(h)]
        skip = {gcol, *ben_cols}
        cat_cols = [j for j, h in enumerate(headers) if h and j not in skip and CATEGORY_HEADER.search(h)]
        typ_cols = [j for j, h in enumerate(headers) if h and j not in skip | set(cat_cols) and TYPE_HEADER.search(h)]
        desc_cols = [j for j, h in enumerate(headers) if h and j not in skip | set(cat_cols) | set(typ_cols) and DESC_HEADER.search(h)]
        mapping.append({"sheet": name, "header_row": hi + 1, "gtin_column": headers[gcol] or f"col{gcol + 1}",
                        "gtin_detected_by": how, "benefit_columns": [headers[j] for j in ben_cols],
                        "category_columns": [headers[j] for j in cat_cols], "type_columns": [headers[j] for j in typ_cols],
                        "description_columns": [headers[j] for j in desc_cols]})
        for off, r in enumerate(body):
            raw = cell_text(r[gcol]) if gcol < len(r) else ""
            if not raw:
                continue
            digits = re.sub(r"[\s\-]", "", raw)
            xl_row = hi + 2 + off  # approximate Excel row (blank rows were dropped); sheet+gtin stays the identity
            key, status, note = gtin_info(digits)
            if "e" in raw.lower() and not digits.isdigit():
                key, status, note = None, "GTIN_INVALID", "scientific notation / text: identifier corrupted by Excel"
            facts, benefits_raw, cat_v, typ_v, desc_v = {}, [], [], [], {}
            for j, h in enumerate(headers):
                if j == gcol or j >= len(r):
                    continue
                t = cell_text(r[j])
                if not t:
                    continue
                if j in ben_cols:
                    benefits_raw.append(t)
                elif j in cat_cols:
                    cat_v.append(t)
                elif j in typ_cols:
                    typ_v.append(t)
                elif j in desc_cols:
                    desc_v[h] = t
                else:
                    facts[h or f"col{j + 1}"] = t
            benefits = []
            for t in benefits_raw:
                if t.strip().lower() not in EMPTY_BENEFIT:
                    benefits += split_benefits(t)
            category, ptype = " / ".join(cat_v) or None, " / ".join(typ_v) or None
            cls = "FROM_FILE_CONFIRM_AT_C1" if category and ptype else "PARTIAL_ASK" if category or ptype else "MISSING_ASK"
            rec = {"sheet": name, "row_hint": xl_row, "gtin_raw": digits, "key": key, "gtin_status": status,
                   "note": note, "facts": facts, "benefits": benefits, "category": category, "product_type": ptype,
                   "classification_status": cls, "existing_description": desc_v,
                   "benefits_status": "PROVIDED" if benefits else "MISSING_DRAFT_FROM_TTX" if facts else "MISSING_NO_TTX"}
            seen_rows.append(rec)
            if key is None:
                issues.append({"issue": "ROW_GTIN_INVALID", "gtin_raw": digits, "detail": note, "sheet": name})
                continue
            if status == "GTIN_INVALID":
                issues.append({"issue": "ROW_GTIN_CHECK_DIGIT", "gtin_raw": digits, "detail": note, "sheet": name})
            rows_by_key.setdefault(key, []).append(rec)

    matched, used_keys = [], set()
    img_by_key = {}
    for im in images:
        if im["key"] is None:
            issues.append({"issue": "IMAGE_GTIN_INVALID", "path": im["path"], "detail": im["note"]})
            continue
        if im["gtin_status"] == "GTIN_INVALID":
            issues.append({"issue": "IMAGE_GTIN_CHECK_DIGIT", "path": im["path"], "detail": im["note"]})
        img_by_key.setdefault(im["key"], []).append(im)

    images_without_row = []
    for key, ims in sorted(img_by_key.items()):
        rows = rows_by_key.get(key)
        if not rows:
            images_without_row += [im["path"] for im in ims]
            issues.append({"issue": "IMAGE_WITHOUT_ROW", "gtin_raw": ims[0]["gtin_raw"], "detail": [im["path"] for im in ims]})
            continue
        used_keys.add(key)
        if len(rows) > 1:
            issues.append({"issue": "DUPLICATE_ROWS_FOR_GTIN", "gtin_raw": ims[0]["gtin_raw"],
                           "detail": f"{len(rows)} rows; first one used, ask the user which is right"})
        flags = sorted({s for s in [rows[0]["gtin_status"]] + [im["gtin_status"] for im in ims] if s != "GTIN_VALID"})
        matched.append({"key": key, "gtin": rows[0]["gtin_raw"], "flags": flags, "images": ims, "row": rows[0],
                        "extra_rows": rows[1:]})
    rows_without_image = [{"gtin": rs[0]["gtin_raw"], "sheet": rs[0]["sheet"], "row_hint": rs[0]["row_hint"]}
                          for k, rs in sorted(rows_by_key.items()) if k not in used_keys]
    for r in rows_without_image:
        issues.append({"issue": "ROW_WITHOUT_IMAGE", "gtin_raw": r["gtin"], "detail": f"{r['sheet']} row ~{r['row_hint']}"})

    groups = {}
    for m in matched:
        r = m["row"]
        g = groups.setdefault((r["category"], r["product_type"]), {"category": r["category"], "product_type": r["product_type"],
                                                                  "status": r["classification_status"], "gtins": []})
        g["gtins"].append(m["gtin"])
    classification_groups = list(groups.values())
    n_cls_ask = sum(1 for m in matched if m["row"]["classification_status"] != "FROM_FILE_CONFIRM_AT_C1")
    n_missing = sum(1 for m in matched if m["row"]["benefits_status"].startswith("MISSING"))
    result = {
        "summary": {"images": len(images), "xlsx_rows": len(seen_rows), "matched_products": len(matched),
                    "benefits_provided": len(matched) - n_missing, "benefits_missing_to_draft": n_missing,
                    "images_without_row": len(images_without_row), "rows_without_image": len(rows_without_image),
                    "issues": len(issues), "classification_to_ask": n_cls_ask,
                    "category_type_groups": len(classification_groups)},
        "classification_groups": classification_groups, "mapping": mapping, "matched": matched, "images_without_row": images_without_row,
        "rows_without_image": rows_without_image, "issues": issues,
    }
    if a.out:
        Path(a.out).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    s = result["summary"]
    print(f"images={s['images']} rows={s['xlsx_rows']} matched={s['matched_products']} "
          f"benefits provided={s['benefits_provided']} / to draft from TTX={s['benefits_missing_to_draft']}")
    for m in mapping:
        print(f"  sheet '{m['sheet']}': GTIN column '{m['gtin_column']}' ({m['gtin_detected_by']}), "
              f"benefit columns: {m['benefit_columns'] or 'none'}, "
              f"category: {m['category_columns'] or 'none'}, type: {m['type_columns'] or 'none'}, "
              f"description: {m['description_columns'] or 'none'}")
    print(f"  category/type groups: {s['category_type_groups']}; products needing a category/type question: {s['classification_to_ask']}")
    for i in issues:
        print(f"  ! {i['issue']}: {i.get('gtin_raw') or i.get('path', '')} {i.get('detail') or ''}")
    if not a.out:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    return 1 if issues else 0


if __name__ == "__main__":
    sys.exit(main())
