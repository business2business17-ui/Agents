#!/usr/bin/env python3
"""Build the Amazon creative production workbook from a JSON plan.

Usage:
  build_workbook.py PLAN.json OUT.xlsx [--template TEMPLATE.xlsx]
  build_workbook.py --example            # print a minimal valid plan

Plan format (all keys optional except assets):
{
  "settings": {"Category": "...", "Mode": "SMART", ...},      # PROJECT_SETTINGS rows
  "assets":   [ {"Product Row ID": "...", "Headline": "...", ...}, ... ],   # ASSET_PLAN, one row = one asset
  "ttx":      [ {...} ],  "content": [ {...} ],  "localization": [ {...} ]   # other sheets
}
Keys must equal the template column headers (references/xlsx-output.md). Unknown keys are
reported and ignored. The template's sample rows are removed; the template itself and the
user's source workbook are never modified. The script appends rows that are not
designer-ready to ISSUES so nothing is silently incomplete.
"""
import argparse
import json
import os
import sys

try:
    import openpyxl
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TEMPLATE = os.path.join(HERE, "..", "assets", "amazon_content_plan_template.xlsx")

SHEETS = {"assets": "ASSET_PLAN", "ttx": "PRODUCT_TTX", "content": "CONTENT_INTELLIGENCE",
          "localization": "LOCALIZATION"}

# Fields that must be non-empty for a row to be designer-ready (references/designer-brief.md section 11).
DESIGNER_READY = ["Placement", "Dimensions", "Requirement Status", "Layout Scheme / Verbal Wireframe",
                  "Product Zone", "Reading Order", "Do Not Cover Areas", "Product Fidelity Lock"]
# Fields that depend on copy being present for non-MAIN image assets.
COPY_FIELDS = ["Headline"]

EXAMPLE = {
    "settings": {"Category": "Haircare", "Product Type": "Shampoo", "Brand": "ExampleBrand",
                 "Target Marketplaces": "DE", "Approval Gate": "SMART (single checkpoint C1)"},
    "assets": [
        {"Project ID": "P1", "Product Row ID": "PRODUCT-001", "Marketplace": "DE", "Language": "de",
         "Category (Row Level)": "Haircare", "Category Status / Confidence": "CONFIRMED",
         "Product Type (Row Level)": "Shampoo", "Product Type Status / Confidence": "CONFIRMED",
         "Content Type": "Listing Image", "Placement": "PDP Carousel", "Asset No.": 1,
         "Module / Surface": "MAIN", "Dimensions": "4000 x 4000", "Aspect Ratio": "1:1",
         "Requirement Status": "PRODUCTION_PRESET (size); AMAZON_REQUIRED (white bg, no text, fill >=85%)",
         "Product Fidelity Lock": "ENABLED - protected immutable product layer",
         "Layout Scheme / Verbal Wireframe": "[CENTER: protected product, 88% of canvas] on pure white",
         "Product Zone": "center", "Reading Order": "product only", "Do Not Cover Areas": "entire product",
         "Source Image / File / URL / Drive Link": "DESCRIPTION ONLY - ASSET NOT CREATED"}
    ],
}


def read_plan(path):
    with open(path, encoding="utf-8") as f:
        plan = json.load(f)
    if not isinstance(plan, dict) or not isinstance(plan.get("assets"), list) or not plan["assets"]:
        sys.exit("plan must be an object with a non-empty 'assets' list")
    return plan


def header_map(ws):
    return {c.value: i + 1 for i, c in enumerate(ws[1]) if c.value}


def clear_rows(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def write_rows(ws, rows, unknown, label):
    hm = header_map(ws)
    for r, row in enumerate(rows, start=2):
        for key, val in row.items():
            col = hm.get(key)
            if col is None:
                unknown.add((label, key))
                continue
            cell = ws.cell(r, col, val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    for name, col in hm.items():
        width = min(60, max(12, len(str(name)) + 2))
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.auto_filter.ref = ws.dimensions


def empty(v):
    return v is None or str(v).strip() == ""


def readiness_issues(assets):
    issues = []
    for i, a in enumerate(assets, start=1):
        label = f"{a.get('Product Row ID', '?')} / {a.get('Placement', '?')} #{a.get('Asset No.', i)}"
        missing = [f for f in DESIGNER_READY if empty(a.get(f))]
        is_main = str(a.get("Module / Surface", "")).strip().upper() == "MAIN"
        if not is_main:
            missing += [f for f in COPY_FIELDS if empty(a.get(f))]
        if missing:
            issues.append({"Asset / Placement": label, "Severity": "HIGH",
                           "Issue": "Not designer-ready: missing " + ", ".join(missing),
                           "Required Correction": "Fill the listed fields (concrete layout, zones, copy) before issuing the designer brief.",
                           "Owner": "Agent", "Status": "OPEN"})
        if str(a.get("Layout Approval Status", "")).upper() not in ("APPROVED", "DELEGATED"):
            issues.append({"Asset / Placement": label, "Severity": "INFO",
                           "Issue": "Layout approval not recorded",
                           "Required Correction": "Set Layout Approval Status to APPROVED (checkpoint C1) or DELEGATED.",
                           "Owner": "Agent", "Status": "OPEN"})
    return issues


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("plan", nargs="?")
    ap.add_argument("out", nargs="?")
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    ap.add_argument("--example", action="store_true")
    a = ap.parse_args()

    if a.example:
        print(json.dumps(EXAMPLE, ensure_ascii=False, indent=2))
        return 0
    if not a.plan or not a.out:
        ap.error("PLAN.json and OUT.xlsx are required")
    if os.path.abspath(a.out) == os.path.abspath(a.template):
        sys.exit("refusing to overwrite the template")
    plan = read_plan(a.plan)
    wb = openpyxl.load_workbook(a.template)
    unknown = set()

    for key, sheet in SHEETS.items():
        rows = plan.get(key) or []
        ws = wb[sheet]
        clear_rows(ws)
        write_rows(ws, rows, unknown, sheet)

    ws = wb["PROJECT_SETTINGS"]
    settings = plan.get("settings") or {}
    existing = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    for k, v in settings.items():
        if k in existing:
            ws.cell(existing[k], 2, v)
        else:
            ws.append([k, v, "ADDED BY PLAN"])

    ws = wb["ISSUES"]
    clear_rows(ws)
    issues = (plan.get("issues") or []) + readiness_issues(plan["assets"])
    for n, issue in enumerate(issues, start=1):
        row = dict(issue)
        row.setdefault("Issue ID", f"ISS-{n:03d}")
        hm = header_map(ws)
        for key, val in row.items():
            if key in hm:
                ws.cell(n + 1, hm[key], val).alignment = Alignment(wrap_text=True, vertical="top")
            else:
                unknown.add(("ISSUES", key))
    ws.freeze_panes = "A2"

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    wb.save(a.out)
    print(f"wrote {a.out}: {len(plan['assets'])} asset rows, {len(issues)} issue rows")
    if unknown:
        print("WARNING unknown columns ignored:", ", ".join(f"{s}:{k}" for s, k in sorted(unknown)), file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
